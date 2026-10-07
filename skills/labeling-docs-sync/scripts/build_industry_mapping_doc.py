"""Build the "Nova - Industry Mapping" workbook (DAT-63) from the live industry mappings.

Five tabs: the source-to-GDU edge table in the original sheet's column order, the GDU industry
catalog with feeder counts and the market/capability kind, the consulting rule, the Fortune 500
headcount vote with its domain allowlist, and an edge-level delta against the Aug 27 original (assets/industry_mapping_original_2026-08-27.csv).

Run (from the main clone so the venv has the project deps; point GDU_REPO at a clean main checkout):
  GDU_REPO=<checkout> uv run --with openpyxl python3 <this file> <out.xlsx>
"""

import csv
import os
import subprocess
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

REPO = Path(os.environ.get("GDU_REPO", Path.home() / "Documents/git/data-universe-pipelines"))
ASSETS = Path(__file__).resolve().parent.parent / "assets"
sys.path.insert(0, str(REPO / "src"))

from constants.industry_mappings import (  # noqa: E402
    CAPABILITY_INDUSTRIES,
    CONSULTING_LINKEDIN_INDUSTRIES,
    CRUNCHBASE_INVESTOR_TO_UNIVERSE_INDUSTRIES,
    CRUNCHBASE_TO_UNIVERSE_INDUSTRIES,
    LINKEDIN_TO_UNIVERSE_INDUSTRIES,
    MARKET_INDUSTRIES,
    DataUniverseIndustries,
)
from constants.organization_constants import (  # noqa: E402
    INDUSTRY_HEADCOUNT_VOTE_DOMAIN_ALLOWLIST,
    INDUSTRY_HEADCOUNT_VOTE_THRESHOLD,
)

HEADER_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
SOURCES = [
    ("Crunchbase", CRUNCHBASE_TO_UNIVERSE_INDUSTRIES),
    ("Crunchbase (Investor)", CRUNCHBASE_INVESTOR_TO_UNIVERSE_INDUSTRIES),
    ("LinkedIn", LINKEDIN_TO_UNIVERSE_INDUSTRIES),
]
UNMAPPED = "(none: source industry carries no GDU industry)"
KIND = {**{i.value: "Market" for i in MARKET_INDUSTRIES}, **{i.value: "Capability" for i in CAPABILITY_INDUSTRIES}}


def add_table(ws, start_row, headers, rows, widths):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h).font = HEADER_FONT
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + 1 + len(rows)


def intro(ws, text, cols=4, height=95):
    ws.cell(row=1, column=1, value=text).alignment = WRAP
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.row_dimensions[1].height = height
    return 3


def edges():
    """(source, source_industry, gdu_industry) in sheet order; unmapped keys yield one UNMAPPED row."""
    for source, mapping in SOURCES:
        for key in sorted(mapping, key=str.lower):
            targets = mapping[key] or [UNMAPPED]
            for gdu in targets:
                yield source, key, gdu.value if hasattr(gdu, "value") else str(gdu)


def main(out):
    wb = Workbook()
    head = subprocess.run(["git", "-C", str(REPO), "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()
    stamp = f" Generated from the mapping code as of {date.today().isoformat()} (commit {head})."

    # --- Industry Mapping ---
    ws = wb.active
    ws.title = "Industry Mapping"
    r = intro(ws, "Source industries from Crunchbase and LinkedIn mapped to GDU's standardized industry taxonomy. One "
                  "source industry can map to several GDU industries, so a source appears once per GDU industry it "
                  "feeds. Crunchbase (Investor) rows are the investor-type field Crunchbase carries on funds and "
                  "investors. A source industry mapped to none was reviewed and deliberately left with no GDU industry. "
                  "Kind says whether the GDU industry names a market the organization sells into or a capability, "
                  "service line or organization type; the consulting rule on its own tab depends on that split." + stamp)
    live = list(edges())
    rows = [[s, k, g, KIND.get(g, "")] for s, k, g in live]
    add_table(ws, r, ["Data Source", "Source Industry", "GDU Industry", "Kind"], rows, [22, 48, 44, 14])
    per_source = Counter(s for s, _, _ in live)
    unmapped = [(s, k) for s, k, g in live if g == UNMAPPED]
    print(f"Edges: {len(live)} {dict(per_source)}; unmapped source industries: {len(unmapped)}")

    # --- GDU Industries ---
    ws = wb.create_sheet("GDU Industries")
    r = intro(ws, "Every GDU industry an organization can carry, with how many source industries feed it from each "
                  "source. Market industries are sectors an organization sells into; capability industries describe "
                  "the organization itself: what it does, what technology it works in, what kind of investor or "
                  "fund it is. Every industry is declared as one or the other.")
    feeders = {src: Counter() for src, _ in SOURCES}
    for s, _, g in live:
        feeders[s][g] += 1
    ind_rows = [[i.value, KIND.get(i.value, ""), feeders["Crunchbase"][i.value], feeders["LinkedIn"][i.value],
                 feeders["Crunchbase (Investor)"][i.value]]
                for i in sorted(DataUniverseIndustries, key=lambda i: i.value.lower())]
    add_table(ws, r, ["GDU Industry", "Kind", "Crunchbase feeders", "LinkedIn feeders", "Investor-type feeders"],
              ind_rows, [48, 14, 20, 18, 20])
    orphans = [i for i, *_ in ind_rows if not any(feeders[s][i] for s, _ in SOURCES)]
    unkinded = [i for i, kind, *_ in ind_rows if not kind]
    print(f"GDU industries: {len(ind_rows)}; market {sum(1 for _, k, *_ in ind_rows if k == 'Market')}, "
          f"capability {sum(1 for _, k, *_ in ind_rows if k == 'Capability')}; "
          f"without a kind: {unkinded or 'none'}; with no feeder: {orphans or 'none'}")

    # --- Consulting rule ---
    ws = wb.create_sheet("Consulting rule")
    r = intro(ws, "A LinkedIn page that declares one of the industries below is saying the organization sells "
                  "consulting, not naming a market it operates in. Such an organization keeps its capability "
                  "industries and drops the market industries it accumulated from the sectors it advises. Only "
                  "LinkedIn triggers this: a LinkedIn page states one or two industries about itself, while a "
                  "Crunchbase record carries an unranked tag list, so a lone Consulting tag says as much about a "
                  "carmaker as about a consultancy.", cols=2)
    add_table(ws, r, ["LinkedIn industry that declares consulting", "GDU industries it maps to"],
              [[k, ", ".join(LINKEDIN_TO_UNIVERSE_INDUSTRIES.get(k, [])) or UNMAPPED] for k in sorted(CONSULTING_LINKEDIN_INDUSTRIES)],
              [48, 70])

    # --- Fortune 500 rule ---
    ws = wb.create_sheet("Fortune 500 rule")
    pct = f"{INDUSTRY_HEADCOUNT_VOTE_THRESHOLD:.0%}"
    r = intro(ws, "The other tabs describe how most organizations get their industries: every industry any of their "
                  "Crunchbase or LinkedIn records maps to. Fortune 500 companies are labeled differently. A large "
                  "company's records include small pages for divisions and subsidiaries, and under the union a "
                  "fifty-person \"CBRE Healthcare\" page puts Healthcare on CBRE. For these organizations the "
                  "records vote by headcount instead:\n"
                  "1. A record with no headcount, or whose industries map to no GDU industry, does not vote.\n"
                  "2. A record's weight is its reported headcount, or the midpoint of its range. A Crunchbase record "
                  "in the open-ended 10,000+ band takes the largest weight among the other voting records, since the "
                  "band hides the true size.\n"
                  "3. Each record gives its full weight to every GDU industry its source industries map to.\n"
                  f"4. An industry stays when its total reaches {pct} of the top industry's total.\n"
                  "5. The consulting rule then runs on what survives.\n"
                  "When no record can vote, the organization carries no industries; there is no fallback to the union.\n\n"
                  "At this point the rule is narrowly scoped to Fortune 500 member companies. It applies only to "
                  "organizations with a domain on the list below: the domains tied to each company on the 2026 "
                  "Fortune 500, plus four pilot companies outside it (openai.com, stripe.com, sodexo.com, "
                  "credit-suisse.com). Any organization sharing one of these domains, such as a single hotel on "
                  "marriott.com, votes the same way. Every other organization keeps the union.", cols=2, height=330)
    add_table(ws, r, ["Domain"], [[d] for d in sorted(INDUSTRY_HEADCOUNT_VOTE_DOMAIN_ALLOWLIST)], [48, 70])
    print(f"Fortune 500 rule: {len(INDUSTRY_HEADCOUNT_VOTE_DOMAIN_ALLOWLIST)} domains, threshold {pct}")

    # --- Changes since original ---
    ws = wb.create_sheet("Changes since original")
    r = intro(ws, "Edge-level differences between the Aug 27 original and this sheet. Removed means the source "
                  "industry no longer maps to that GDU industry; added means it does now and did not on 27 Aug. "
                  "Missing from original marks edges the mapping already held on 27 Aug that the original sheet "
                  "left out, most of them LinkedIn. Unchanged edges are not listed.")
    with open(ASSETS / "industry_mapping_original_2026-08-27.csv") as fh:
        original = {(row["source"], row["source_industry"], row["gdu_industry"]) for row in csv.DictReader(fh)}
    # What the code held the day the original sheet was made; edges here but not in the sheet
    # were never exported, not added since.
    with open(ASSETS / "industry_mapping_code_2026-08-27.csv") as fh:
        code_then = {(row["source"], row["source_industry"], row["gdu_industry"]) for row in csv.DictReader(fh)}
    current = {(s, k, g) for s, k, g in live if g != UNMAPPED}
    orig_keys = {(s, k) for s, k, _ in original}
    cur_keys = {(s, k) for s, k, _ in current} | {(s, k) for s, k in unmapped}
    ch_rows = []
    for s, k, g in sorted(original - current, key=lambda e: (e[0], e[1].lower(), e[2])):
        note = "Source industry now carries no GDU industry" if (s, k) in {(a, b) for a, b in unmapped} else \
               ("Source industry no longer in the mapping" if (s, k) not in cur_keys else "")
        ch_rows.append([s, k, g, "Removed", note])
    for s, k, g in sorted(current - original, key=lambda e: (e[0], e[1].lower(), e[2])):
        if (s, k, g) in code_then:
            change, note = "Missing from original", "Was in the mapping on 27 Aug but not in the original sheet"
        else:
            change = "Added"
            note = "New GDU industry" if not any(og == g for *_, og in original) else \
                   ("New source industry" if (s, k) not in orig_keys else "")
        ch_rows.append([s, k, g, change, note])
    add_table(ws, r, ["Data Source", "Source Industry", "GDU Industry", "Change", "Notes"], ch_rows, [22, 48, 44, 12, 46])
    tally = Counter(row[3] for row in ch_rows)
    new_labels = sorted({g for _, _, g in current} - {g for *_, g in original})
    gone_labels = sorted({g for *_, g in original} - {g for _, _, g in current})
    print(f"Original edges: {len(original)}; unchanged {len(original & current)}; {dict(tally)}; "
          f"new GDU industries: {new_labels}; GDU industries with no feeder any more: {gone_labels}")

    dirty = subprocess.run(["git", "-C", str(REPO), "status", "--porcelain", "--", "src/constants/industry_mappings.py"],
                           capture_output=True, text=True).stdout
    print(f"Built from {head}{' (DIRTY source file)' if dirty else ''}")
    wb.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1]))

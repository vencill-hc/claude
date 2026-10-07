"""Build the "Nova - Function Mapping" workbook (DAT-64) from the live job-function spec.

Four tabs: the keyword table in match order, the functions catalog, the removed
keywords, and a delta against the Aug 27 original sheet (assets/function_mapping_original_2026-08-27.csv).
Keyword rationale and function descriptions come from docs/job_function_taxonomy.md
in the same checkout, so the workbook cannot say something the repo does not.

Run (from a clean checkout of main, or a worktree of it):
  GDU_REPO=<checkout> uv run --with openpyxl python3 <this file> <out.xlsx>
"""

import csv
import os
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

REPO = Path(os.environ.get("GDU_REPO", Path.home() / "Documents/git/data-universe-pipelines"))
ASSETS = Path(__file__).resolve().parent.parent / "assets"
sys.path.insert(0, str(REPO / "src"))

from data_types.job_types import LIVE_JOB_FUNCTIONS, JobFunction  # noqa: E402
from utils.employment_utils import JobFunctionParseFailure, JobFunctionParser  # noqa: E402
from utils.job_function_taxonomy import REMOVED_KEYWORDS, SPEC_BANDS, RemovalReason  # noqa: E402
from utils.string_validator_utils import JOB_TITLE_NORMALIZATION_SYNONYMS  # noqa: E402

HEADER_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

BAND_LABEL = {
    1: "1 - carve-out",
    2: "2 - specific phrase",
    3: "3 - single word",
    4: "4 - backstop",
    5: "5 - frontline vocabulary",
}
REMOVAL_LABEL = {
    RemovalReason.OVER_CAPTURE: "Matched too broadly",
    RemovalReason.OFF_AXIS: "Not a job function (level, engagement, or industry)",
}
INTRO = (
    "GDU reads a person's job function from their job title. The title is normalized first: lowercased, "
    "punctuation stripped, common abbreviations unified (chief executive officer becomes ceo). It is then "
    "scanned against the keywords below in the order listed, and the first keyword found decides the function. "
    "Keywords are matched as prefixes, so engineer also catches engineering and engineers. A keyword written "
    "with underscores, like _it_ or _sales_, is an exact-word sentinel: the title pre-processing wraps that "
    "word when it appears on its own, so _it_ catches 'IT manager' but not 'Italian' or 'iteration'. "
    "The five priority bands: carve-outs come "
    "first so software engineer reads as Software engineering before the bare word engineer can fire, then "
    "specific phrases, then single words, then short backstop terms, then frontline vocabulary added below "
    "everything that already existed. A keyword marked BLOCKER stops the scan with no function assigned."
)


def add_table(ws, start_row, headers, rows, widths):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h).font = HEADER_FONT
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + 1 + len(rows)


def intro(ws, text):
    ws.cell(row=1, column=1, value=text).alignment = WRAP
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.row_dimensions[1].height = 110
    return 3


def scrub(text: str) -> str:
    """Reader-facing text: sentinels stay (the intro explains them), regex fragments and code names go."""
    text = text.replace("`", "")
    text = re.sub(r"\\[bw]\*?", "", text)  # regex fragments like \blean\w*\b read as the bare word
    return text.replace("_perform_subs", "title pre-processing").replace("normalize_job_title", "title normalization")


def display(kw: str) -> tuple[str, str]:
    """Sentinel `_x_` keys match the exact word; everything else is a prefix."""
    if kw == "_assistant_":
        return kw, "Only when the whole title is the single word assistant"
    if kw == "cxo":
        return kw, "Any chief-something-officer title not covered by a more specific key"
    if "_" not in kw:
        return kw, ""
    note = "Exact-word sentinel"
    if "_sales_" in kw:
        note += "; _sales_ also stands in for sales engineer, solution engineer and partner manager"
    if "_chef_" in kw:
        note += "; chef and chefs only, so chefe and cheffe (Portuguese and French for boss) stay out"
    return kw, note


def parse_doc(path: Path):
    """Pull the functions table and the keyword-rulings table out of the markdown doc."""
    functions, rulings, section = {}, {}, None
    for line in path.read_text().splitlines():
        if line.startswith("## "):
            section = line[3:].strip()
            continue
        if not line.startswith("| ") or line.startswith("| Function |") or line.startswith("| Band |") or line.startswith("|---"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split(" | ")]
        if section == "Functions" and len(cells) >= 4:
            functions[cells[0]] = cells[3]
        elif section == "Keyword rulings" and len(cells) >= 4:
            rulings[cells[1].strip("`")] = scrub(cells[3])
    return functions, rulings


def main(out: Path):
    functions_desc, rulings = parse_doc(REPO / "docs/job_function_taxonomy.md")
    wb = Workbook()

    # --- Job Function Mapping ---
    ws = wb.active
    ws.title = "Job Function Mapping"
    r = intro(ws, INTRO)
    rows, live_spec = [], {}
    for band_no, (_, band) in enumerate(SPEC_BANDS, 1):
        for kw, fn in band.items():
            shown, note = display(kw)
            why = rulings.get(kw, "")
            notes = "; ".join(p for p in (note, why) if p)
            target = fn.value if fn is not None else "BLOCKER - no function assigned"
            rows.append([BAND_LABEL[band_no], shown, target, notes])
            live_spec[kw] = fn
    add_table(ws, r, ["Match Priority", "Keyword in Job Title", "Assigned Job Function", "Notes"], rows, [22, 34, 36, 90])
    print(f"Keyword rows: {len(rows)} across {len(SPEC_BANDS)} bands; blockers: {sum(fn is None for fn in live_spec.values())}")

    # --- Functions ---
    ws = wb.create_sheet("Functions")
    r = intro(ws, "Every job function a title can resolve to, grouped, with what it covers. The group is a reading "
                  "aid only: the classifier assigns the function and stores nothing above it. Functions retired since "
                  "the original sheet are absent here; a person whose stored function was retired is relabeled from "
                  "their title.")
    # Groups left the code on 2026-09-10 (the taxonomy is flat); this curated map is the
    # last one that existed plus the functions added since. Keyed by enum name so renames don't break it.
    with open(ASSETS / "function_groups.csv") as fh:
        groups = {JobFunction[row["function_enum"]].value: row["group"] for row in csv.DictReader(fh)}
    fn_rows = [[groups.get(f.value, ""), f.value, functions_desc.get(f.value, "")]
               for f in sorted(LIVE_JOB_FUNCTIONS, key=lambda f: (groups.get(f.value, "~").lower(), f.value.lower()))]
    missing = [f for _, f, d in fn_rows if not d]
    ungrouped = [f for g, f, _ in fn_rows if not g]
    add_table(ws, r, ["Group", "Job Function", "What it covers"], fn_rows, [34, 40, 110])
    print(f"Functions: {len(fn_rows)} in {len(set(groups.values()))} groups; without a description: {missing or 'none'}; "
          f"without a group: {ungrouped or 'none'}")

    # --- Removed Keywords ---
    ws = wb.create_sheet("Removed Keywords")
    r = intro(ws, "Keywords that were in the original sheet or a later revision and have since been deleted. "
                  "A title that only ever matched one of these now carries no function.")
    rm_rows = [[display(kw)[0], REMOVAL_LABEL[reason], rulings.get(kw, "")] for kw, reason in REMOVED_KEYWORDS.items()]
    add_table(ws, r, ["Keyword", "Why it was removed", "Notes"], rm_rows, [30, 50, 100])
    print(f"Removed keywords: {len(rm_rows)}")

    # --- Changes since original ---
    ws = wb.create_sheet("Changes since original")
    r = intro(ws, "Every keyword in the Aug 27 original, with where it resolves now. 'Function now' is what the "
                  "live parser returns when the keyword is run through it as a title on its own. Still present "
                  "means the keyword is in the table above; subsumed means a broader keyword now catches it; "
                  "removed means the keyword was deleted on purpose and the title gets no function.")
    with open(ASSETS / "function_mapping_original_2026-08-27.csv") as fh:
        original = [(row["keyword"], row["function"]) for row in csv.DictReader(fh)]
    ch_rows, tally, untracked = [], {}, []
    for shown_kw, old_fn in original:
        # The original key may now live under the abbreviation normalization
        # applies first (chief executive officer -> ceo), or as its exact-word sentinel.
        kw = JOB_TITLE_NORMALIZATION_SYNONYMS.get(shown_kw, shown_kw)
        kw = f"_{kw}_" if kw not in live_spec and f"_{kw}_" in live_spec else kw
        outcome = JobFunctionParser.parse_outcome(shown_kw.replace("_", ""))
        if isinstance(outcome, JobFunctionParseFailure):
            # REMOVED_KEY covers both a live blocker and a deleted off-axis key; only the former is a blocker.
            new_v = "BLOCKER - no function assigned" if kw in live_spec and live_spec[kw] is None else ""
        else:
            new_v = outcome.value
        if kw in live_spec:
            status = "Still present" + ("" if new_v == old_fn else ", re-pointed")
        elif kw in REMOVED_KEYWORDS:
            status = f"Removed: {REMOVAL_LABEL[REMOVED_KEYWORDS[kw]]}"
        elif new_v:
            status = "Subsumed by a broader keyword"
        else:
            status = "Dropped with no removal record"
            untracked.append(shown_kw)
        key = re.split(r"[:,]", status)[0]
        tally[key] = tally.get(key, 0) + 1
        ch_rows.append([shown_kw, old_fn, new_v, status, rulings.get(kw, "")])
    add_table(ws, r, ["Keyword", "Function in original", "Function now", "Status", "Notes"], ch_rows, [30, 30, 36, 34, 90])
    print(f"Original keywords: {len(original)}; {tally}")
    if untracked:
        print(f"WARNING original keywords gone from the spec with no REMOVED_KEYWORDS entry: {untracked}")

    head = subprocess.run(["git", "-C", str(REPO), "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()
    dirty = subprocess.run(["git", "-C", str(REPO), "status", "--porcelain", "--", "src/utils/job_function_taxonomy.py",
                            "src/data_types/job_types.py", "docs/job_function_taxonomy.md"], capture_output=True, text=True).stdout
    print(f"Built from {head}{' (DIRTY source files)' if dirty else ''}")
    wb.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1]))

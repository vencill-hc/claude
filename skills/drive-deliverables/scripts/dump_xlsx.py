# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""Dump every sheet of an xlsx as pipe-delimited rows (throwaway)."""
import sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
for ws in wb.worksheets:
    print(f"\n===== SHEET: {ws.title} ({ws.max_row}x{ws.max_column})")
    for row in ws.iter_rows(values_only=True):
        if not any(c not in (None, "") for c in row):
            continue
        cells = ["" if c is None else str(c).replace("\n", " / ") for c in row]
        while cells and cells[-1] == "":
            cells.pop()
        print(" | ".join(cells))
